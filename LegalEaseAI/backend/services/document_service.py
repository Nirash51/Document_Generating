import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF


def sanitize_text(text: str) -> str:

    replacements = {

        "\u2018": "'",
        "\u2019": "'",

        "\u201c": '"',
        "\u201d": '"',

        "\u2013": "-",
        "\u2014": "-",

        "\u2026": "...",

        "\u00a0": " ",
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new
        )

    return "".join(
        character
        for character in text
        if character in "\n\r\t"
        or ord(character) >= 32
    )


def terms_from_text(
    terms: str
) -> list[str]:

    return [
        item.strip(" \t-")
        for item in terms.split(";")
        if item.strip()
    ]


def markdown_lines(
    text: str
) -> list[tuple[str, str]]:

    result = []

    clean_text = sanitize_text(text)

    for raw_line in clean_text.splitlines():

        line = raw_line.strip()

        if not line:

            result.append(
                ("blank", "")
            )

        elif line.startswith("### "):

            result.append(
                ("h3", line[4:].strip())
            )

        elif line.startswith("## "):

            result.append(
                ("h2", line[3:].strip())
            )

        elif line.startswith("# "):

            result.append(
                ("h1", line[2:].strip())
            )

        elif re.match(
            r"^[-*]\s+",
            line
        ):

            result.append(
                (
                    "bullet",
                    re.sub(
                        r"^[-*]\s+",
                        "",
                        line
                    )
                )
            )

        elif re.match(
            r"^\d+[.)]\s+",
            line
        ):

            result.append(
                (
                    "number",
                    re.sub(
                        r"^\d+[.)]\s+",
                        "",
                        line
                    )
                )
            )

        else:

            result.append(
                (
                    "paragraph",
                    line
                )
            )

    return result


def format_docx(
    text: str,
    doc_type: str,
    terms: list[str] | None = None,
    brand_name: str = "LegalEase",
    logo_path: str | None = None,
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles

    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    if logo_path and Path(logo_path).exists():

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            paragraph.add_run().add_picture(
                logo_path,
                width=Inches(1.25)
            )

        except Exception:

            pass

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        doc_type.upper()
    )

    run.bold = True

    run.font.name = "Times New Roman"

    run.font.size = Pt(16)

    for kind, content in markdown_lines(text):

        if kind == "blank":
            continue

        if kind == "h1":

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                content
            )

            run.bold = True
            run.font.size = Pt(15)

        elif kind == "h2":

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                content
            )

            run.bold = True
            run.font.size = Pt(13)

        elif kind == "h3":

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                content
            )

            run.bold = True
            run.font.size = Pt(12)

        elif kind == "bullet":

            document.add_paragraph(
                content,
                style="List Bullet"
            )

        elif kind == "number":

            document.add_paragraph(
                content,
                style="List Number"
            )

        else:

            document.add_paragraph(
                content
            )

    if terms:

        document.add_heading(
            "Key Terms",
            level=2
        )

        table = document.add_table(
            rows=1,
            cols=2
        )

        table.style = "Table Grid"

        table.rows[0].cells[0].text = "No."

        table.rows[0].cells[1].text = (
            "Term / Condition"
        )

        for index, term in enumerate(
            terms,
            start=1
        ):

            cells = table.add_row().cells

            cells[0].text = str(index)

            cells[1].text = term

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        f"{brand_name} | AI-assisted draft | "
        "Review before use"
    )

    footer_run.italic = True

    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()


def pdf_safe(text: str) -> str:

    text = sanitize_text(text)

    return (
        text
        .encode(
            "latin-1",
            errors="replace"
        )
        .decode("latin-1")
    )


class LegalEasePDF(FPDF):

    def __init__(
        self,
        brand_name: str
    ):

        super().__init__(
            orientation="P",
            unit="mm",
            format="A4"
        )

        self.brand_name = brand_name

    def header(self):

        self.set_font(
            "Helvetica",
            "B",
            9
        )

        self.cell(
            self.epw,
            8,
            pdf_safe(
                self.brand_name
            ),
            align="C",
            new_x="LMARGIN",
            new_y="NEXT"
        )

        self.ln(3)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "I",
            8
        )

        self.cell(
            self.epw,
            8,
            pdf_safe(
                f"{self.brand_name} | "
                f"AI-assisted draft | "
                f"Page {self.page_no()}"
            ),
            align="C",
            new_x="LMARGIN",
            new_y="TOP"
        )


def format_pdf(
    text: str,
    doc_type: str,
    terms: list[str] | None = None,
    brand_name: str = "LegalEase",
) -> bytes:

    pdf = LegalEasePDF(
        brand_name
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=18
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "B",
        16
    )

    pdf.multi_cell(
        pdf.epw,
        9,
        pdf_safe(
            doc_type.upper()
        ),
        align="C",
        new_x="LMARGIN",
        new_y="NEXT",
    )

    pdf.ln(4)

    for kind, content in markdown_lines(text):

        pdf.set_x(pdf.l_margin)

        if kind == "blank":

            pdf.ln(2)

            continue

        if kind in {
            "h1",
            "h2",
            "h3"
        }:

            size = {
                "h1": 14,
                "h2": 12,
                "h3": 11,
            }[kind]

            pdf.set_font(
                "Helvetica",
                "B",
                size
            )

            pdf.multi_cell(
                pdf.epw,
                7,
                pdf_safe(content),
                new_x="LMARGIN",
                new_y="NEXT",
            )

            pdf.ln(1)

        elif kind == "bullet":

            pdf.set_font(
                "Helvetica",
                "",
                10.5
            )

            pdf.multi_cell(
                pdf.epw,
                6,
                pdf_safe(
                    f"• {content}"
                ),
                new_x="LMARGIN",
                new_y="NEXT",
            )

        elif kind == "number":

            pdf.set_font(
                "Helvetica",
                "",
                10.5
            )

            pdf.multi_cell(
                pdf.epw,
                6,
                pdf_safe(
                    f"• {content}"
                ),
                new_x="LMARGIN",
                new_y="NEXT",
            )

        else:

            pdf.set_font(
                "Helvetica",
                "",
                10.5
            )

            pdf.multi_cell(
                pdf.epw,
                6,
                pdf_safe(content),
                new_x="LMARGIN",
                new_y="NEXT",
            )

    if terms:

        pdf.set_x(pdf.l_margin)
        pdf.ln(3)

        pdf.set_font(
            "Helvetica",
            "B",
            12
        )

        pdf.cell(
            pdf.epw,
            7,
            "Key Terms",
            new_x="LMARGIN",
            new_y="NEXT",
        )

        pdf.ln(2)

        for index, term in enumerate(
            terms,
            start=1
        ):

            pdf.set_x(pdf.l_margin)

            pdf.set_font(
                "Helvetica",
                "",
                10
            )

            pdf.multi_cell(
                pdf.epw,
                6,
                pdf_safe(
                    f"{index}. {term}"
                ),
                new_x="LMARGIN",
                new_y="NEXT",
            )

    return bytes(
        pdf.output()
    )


def format_txt(
    text: str
) -> bytes:

    return sanitize_text(
        text
    ).encode("utf-8")